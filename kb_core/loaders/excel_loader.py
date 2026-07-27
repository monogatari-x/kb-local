from pathlib import Path

from kb_core.exceptions import FilePathError
from kb_core.loaders.base import BaseLoader, Block, LoadedDocument


class ExcelLoader(BaseLoader):
    def supported_extensions(self) -> set[str]:
        return {".xlsx", ".xls"}

    def load(self, path: Path) -> LoadedDocument:
        if not path.exists():
            raise FilePathError(f"Excel file not found: {path}")
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        blocks: list[Block] = []
        all_text: list[str] = []
        sheet_count = len(wb.sheetnames)
        for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            rows_text: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows_text.append("\t".join(cells))
            if rows_text:
                block_text = "\n".join(rows_text)
                blocks.append(
                    Block(
                        text=block_text,
                        start_line=sheet_idx,
                        end_line=sheet_idx,
                        kind="sheet",
                        extra={"sheet": sheet_name},
                    )
                )
                all_text.append(block_text)
        wb.close()
        return LoadedDocument(
            text="\n\n".join(all_text),
            language=None,
            blocks=blocks,
            meta={"sheet_count": sheet_count},
        )
