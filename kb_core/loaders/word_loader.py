from pathlib import Path

from kb_core.exceptions import FilePathError
from kb_core.loaders.base import BaseLoader, Block, LoadedDocument


class WordLoader(BaseLoader):
    def supported_extensions(self) -> set[str]:
        return {".docx"}

    def load(self, path: Path) -> LoadedDocument:
        if not path.exists():
            raise FilePathError(f"Word file not found: {path}")
        from docx import Document

        doc = Document(str(path))
        blocks: list[Block] = []
        all_text: list[str] = []
        para_idx = 0
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            para_idx += 1
            blocks.append(
                Block(
                    text=text,
                    start_line=para_idx,
                    end_line=para_idx,
                    kind="paragraph",
                )
            )
            all_text.append(text)
        return LoadedDocument(
            text="\n".join(all_text),
            language=None,
            blocks=blocks,
            meta={},
        )
